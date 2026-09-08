import asyncio
import csv
import io
import ipaddress
import math
import os
import socket
from collections.abc import AsyncIterator, Iterable
from datetime import datetime, timedelta
from typing import Any
from urllib.parse import urlsplit

import httpx
from fastapi import HTTPException, status
from openpyxl import Workbook
from sqlglot import exp, parse

from .schema import (
    AlarmQueryIn,
    HistoryMode,
    HistoryQueryIn,
    HistoryWriteIn,
    RealDataQueryIn,
    RealDataWriteIn,
    RegisterPointsIn,
    RegisteredQueryIn,
    SqlQueryIn,
)


class ZijinBridgeService:
    """紫金桥跨平台实时数据库 REST API 适配器。

    设计目标：
    1. 不在服务端持久化紫金桥用户名、密码或 token；
    2. 对大量历史数据自动按时间与点位切片，避免单次请求过大；
    3. 默认只允许访问本机/内网 REST 服务，降低把管理端变成 SSRF 代理的风险。
    """

    DEFAULT_TIMEOUT = float(os.getenv("ZIJINBRIDGE_TIMEOUT", "30"))
    ALLOW_PUBLIC_HOSTS = os.getenv("ZIJINBRIDGE_ALLOW_PUBLIC_HOSTS", "0").lower() in {"1", "true", "yes"}
    MAX_EXPORT_ROWS = int(os.getenv("ZIJINBRIDGE_MAX_EXPORT_ROWS", "2000000"))
    MAX_QUERY_CELLS = int(os.getenv("ZIJINBRIDGE_MAX_QUERY_CELLS", "2000000"))

    @classmethod
    async def normalize_base_url(cls, raw_url: str) -> str:
        raw_url = raw_url.strip().rstrip("/")
        parsed = urlsplit(raw_url)
        if parsed.scheme not in {"http", "https"} or not parsed.hostname:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="紫金桥地址必须是 http/https URL")
        if parsed.path not in {"", "/"} or parsed.query or parsed.fragment:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="紫金桥地址只填写协议、主机和端口，不要携带路径/查询参数")
        if parsed.username or parsed.password:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="连接地址中禁止携带用户名或密码")
        if cls.ALLOW_PUBLIC_HOSTS:
            return raw_url

        host = parsed.hostname
        if host.lower() == "localhost":
            return raw_url

        def _is_private(value: str) -> bool:
            try:
                addr = ipaddress.ip_address(value)
            except ValueError:
                return False
            return addr.is_private or addr.is_loopback or addr.is_link_local

        if _is_private(host):
            return raw_url

        try:
            loop = asyncio.get_running_loop()
            infos = await loop.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)
            resolved = {item[4][0] for item in infos}
        except OSError as exc:
            raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=f"无法解析紫金桥主机：{host}") from exc

        if not resolved or not all(_is_private(ip) for ip in resolved):
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="默认仅允许连接本机或内网紫金桥服务；如确需公网地址，请设置 ZIJINBRIDGE_ALLOW_PUBLIC_HOSTS=1",
            )
        return raw_url

    @classmethod
    async def request(
        cls,
        base_url: str,
        method: str,
        path: str,
        *,
        token: str | None = None,
        params: dict[str, Any] | None = None,
        json_body: dict[str, Any] | None = None,
        check_code: bool = True,
    ) -> dict[str, Any]:
        base_url = await cls.normalize_base_url(base_url)
        query = {key: value for key, value in (params or {}).items() if value is not None}
        if token:
            query["token"] = token

        try:
            async with httpx.AsyncClient(base_url=base_url, timeout=cls.DEFAULT_TIMEOUT, follow_redirects=False) as client:
                response = await client.request(method, path, params=query, json=json_body)
        except httpx.TimeoutException as exc:
            raise HTTPException(status_code=status.HTTP_504_GATEWAY_TIMEOUT, detail="连接紫金桥 REST 服务超时") from exc
        except httpx.HTTPError as exc:
            raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=f"连接紫金桥 REST 服务失败：{exc.__class__.__name__}") from exc

        if response.is_redirect:
            raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail="紫金桥 REST 服务返回了重定向，已拒绝自动跟随")
        try:
            payload = response.json()
        except ValueError as exc:
            detail = response.text[:300].replace("\n", " ")
            raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=f"紫金桥返回非 JSON 数据：{detail}") from exc
        if not isinstance(payload, dict):
            raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail="紫金桥返回格式不是 JSON 对象")

        if response.is_error:
            raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=f"紫金桥 HTTP {response.status_code}: {payload.get('message', '')}")
        if check_code and payload.get("code", 0) != 0:
            message = payload.get("message") or payload.get("msg") or "未知错误"
            raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail=f"紫金桥 API code={payload.get('code')}: {message}")
        return payload

    @classmethod
    async def login(cls, base_url: str, user: str, password: str) -> dict[str, Any]:
        # 上游文档要求 user/pass 使用查询参数；此处不记录、不持久化口令。
        return await cls.request(base_url, "GET", "/api/userlogin/", params={"user": user, "pass": password})

    @classmethod
    async def logout(cls, base_url: str, token: str) -> dict[str, Any]:
        return await cls.request(base_url, "GET", "/api/userlogout", token=token)

    @classmethod
    async def status(cls, base_url: str, token: str | None) -> dict[str, Any]:
        normalized = await cls.normalize_base_url(base_url)
        try:
            payload = await cls.request(normalized, "GET", "/api/realdir", token=token, check_code=False)
            code = payload.get("code")
            return {
                "base_url": normalized,
                "reachable": True,
                "authenticated": code in (None, 0),
                "upstream_code": code if isinstance(code, int) else None,
                "message": payload.get("message") or payload.get("msg") or "REST 服务可访问",
            }
        except HTTPException as exc:
            if exc.status_code in {502, 504}:
                return {
                    "base_url": normalized,
                    "reachable": False,
                    "authenticated": None,
                    "upstream_code": None,
                    "message": str(exc.detail),
                }
            raise

    @classmethod
    async def realdir(cls, base_url: str, token: str | None, node_path: str = "") -> dict[str, Any]:
        suffix = "/" + node_path.strip("/") if node_path.strip("/") else ""
        return await cls.request(base_url, "GET", f"/api/realdir{suffix}", token=token)

    @classmethod
    async def node_points(
        cls,
        base_url: str,
        token: str | None,
        node_path: str,
        *,
        recursion: bool = True,
        decimal: int | None = None,
    ) -> dict[str, Any]:
        suffix = node_path.strip("/")
        return await cls.request(
            base_url,
            "GET",
            f"/api/realdata/{suffix}",
            token=token,
            params={"recursion": int(recursion), "valueonly": 0, "decimal": decimal},
        )

    @classmethod
    async def realdata(cls, payload: RealDataQueryIn, token: str | None) -> dict[str, Any]:
        params: dict[str, Any] = {"decimal": payload.decimal, "valueonly": int(payload.value_only)}
        if payload.names:
            params["names"] = ",".join(payload.names)
        else:
            params["tags"] = ",".join(payload.tags or [])
            params["pars"] = ",".join(payload.pars or [])
        return await cls.request(payload.base_url, "GET", "/api/realdata", token=token, params=params)

    @classmethod
    async def register_points(cls, payload: RegisterPointsIn, token: str) -> dict[str, Any]:
        return await cls.request(
            payload.base_url,
            "GET",
            "/api/realdata",
            token=token,
            params={"regname": payload.regname, "names": ",".join(payload.names)},
        )

    @classmethod
    async def registered_realdata(cls, payload: RegisteredQueryIn, token: str) -> dict[str, Any]:
        return await cls.request(
            payload.base_url,
            "GET",
            "/api/realdata",
            token=token,
            params={"regname": payload.regname, "decimal": payload.decimal},
        )

    @classmethod
    async def write_realdata(cls, payload: RealDataWriteIn, token: str | None) -> dict[str, Any]:
        return await cls.request(
            payload.base_url,
            "POST",
            "/api/realdata",
            token=token,
            json_body={"items": [item.model_dump() for item in payload.items]},
        )

    @staticmethod
    def _parse_time(value: str) -> datetime:
        try:
            return datetime.fromisoformat(value)
        except ValueError as exc:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=f"时间格式错误：{value}") from exc

    @classmethod
    def _range_windows(cls, payload: HistoryQueryIn) -> list[tuple[str, str, int]]:
        start = cls._parse_time(payload.start_time or "")
        end = cls._parse_time(payload.end_time or "")
        if end <= start:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="end_time 必须晚于 start_time")
        interval = int(payload.interval or 1)
        delta_ms = (end - start).total_seconds() * 1000
        estimated = int(math.ceil(delta_ms / interval))
        if estimated * len(payload.names) > cls.MAX_QUERY_CELLS:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail=f"查询预计产生 {estimated * len(payload.names):,} 个数据单元，超过管理端安全上限；请增大间隔或缩小时间范围",
            )
        windows: list[tuple[str, str, int]] = []
        cursor = start
        chunk_delta = timedelta(milliseconds=interval * payload.samples_per_request)
        while cursor < end:
            window_end = min(end, cursor + chunk_delta)
            windows.append((cursor.isoformat(timespec="milliseconds"), window_end.isoformat(timespec="milliseconds"), interval))
            cursor = window_end
        return windows

    @staticmethod
    def _chunks(values: list[str], size: int) -> Iterable[list[str]]:
        for index in range(0, len(values), size):
            yield values[index : index + size]

    @classmethod
    async def history(cls, payload: HistoryQueryIn, token: str | None) -> dict[str, Any]:
        merged: dict[str, list[Any]] = {name: [] for name in payload.names}
        request_count = 0

        if payload.mode == HistoryMode.TIMES:
            for query_time in payload.times:
                cls._parse_time(query_time)
            cells = len(payload.times) * len(payload.names)
            if cells > cls.MAX_QUERY_CELLS:
                raise HTTPException(
                    status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                    detail=f"查询预计产生 {cells:,} 个数据单元，超过管理端安全上限；请减少点位或时间点",
                )
            time_batches = list(cls._chunks(payload.times, payload.samples_per_request))
            query_units = [("times", batch) for batch in time_batches]
        else:
            query_units = [("range", window) for window in cls._range_windows(payload)]

        for mode, unit in query_units:
            for name_batch in cls._chunks(payload.names, payload.names_per_request):
                params: dict[str, Any] = {
                    "names": ",".join(name_batch),
                    "decimal": payload.decimal,
                    "valueonly": int(payload.value_only),
                }
                if mode == "times":
                    params["times"] = ",".join(unit)  # type: ignore[arg-type]
                else:
                    start_time, end_time, interval = unit  # type: ignore[misc]
                    params.update({"startTime": start_time, "endTime": end_time, "interval": interval})

                result = await cls.request(payload.base_url, "GET", "/api/hisdata", token=token, params=params)
                request_count += 1
                for item in result.get("items", []):
                    name = item.get("name")
                    if name in merged:
                        merged[name].extend(item.get("vals") or [])

        return {
            "code": 0,
            "mode": payload.mode,
            "items": [{"name": name, "vals": merged[name]} for name in payload.names],
            "meta": {
                "names": len(payload.names),
                "upstream_requests": request_count,
                "value_only": payload.value_only,
                "interval": payload.interval if payload.mode == HistoryMode.RANGE else None,
                "times": payload.times if payload.mode == HistoryMode.TIMES else None,
                "start_time": payload.start_time,
                "end_time": payload.end_time,
            },
        }

    @classmethod
    async def write_history(cls, payload: HistoryWriteIn, token: str | None) -> dict[str, Any]:
        return await cls.request(
            payload.base_url,
            "POST",
            "/api/hisdata",
            token=token,
            json_body={"items": [item.model_dump() for item in payload.items]},
        )

    @classmethod
    async def sql(cls, payload: SqlQueryIn, token: str | None) -> dict[str, Any]:
        try:
            statements = parse(payload.query)
        except Exception as exc:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=f"SQL 解析失败：{exc}") from exc
        if not statements:
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="SQL 不能为空")

        # 管理端默认只开放查询语句，写入请使用明确的实时/历史数据写入 API。
        # SELECT ... INTO 也可能产生写入副作用，因此只读模式显式拒绝。
        read_only_safe = all(
            isinstance(statement, exp.Query) and statement.find(exp.Into) is None
            for statement in statements
        )
        if payload.read_only and not read_only_safe:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="SQL 控制台默认只允许无写入副作用的 SELECT/查询语句")
        if not payload.read_only and os.getenv("ZIJINBRIDGE_ALLOW_SQL_WRITE", "0").lower() not in {"1", "true", "yes"}:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="服务端未启用 ZIJINBRIDGE_ALLOW_SQL_WRITE=1")

        return await cls.request(payload.base_url, "POST", "/api/SQL", token=token, json_body={"query": payload.query})

    @classmethod
    async def alarms(cls, payload: AlarmQueryIn, token: str | None) -> dict[str, Any]:
        params = {
            "realalmcount": payload.realalmcount,
            "org": payload.org,
            "unit": payload.unit,
            "level": payload.level,
            "tags": ",".join(payload.tags) if payload.tags else None,
            "groupbytag": int(payload.group_by_tag),
            "startTime": payload.start_time,
            "endTime": payload.end_time,
        }
        return await cls.request(payload.base_url, "GET", "/api/alarmdata", token=token, params=params)

    @staticmethod
    def _safe_export_value(value: Any) -> Any:
        """避免 CSV/XLSX 中的字符串被表格软件当成公式执行。"""
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    @classmethod
    def _history_rows(cls, payload: HistoryQueryIn, result: dict[str, Any]) -> Iterable[tuple[str, str, Any]]:
        for item in result.get("items", []):
            name = str(item.get("name", ""))
            values = item.get("vals") or []
            for index, value in enumerate(values):
                if isinstance(value, dict):
                    yield name, str(value.get("time", "")), value.get("val")
                    continue
                if payload.mode == HistoryMode.TIMES:
                    row_time = payload.times[index] if index < len(payload.times) else ""
                else:
                    start = cls._parse_time(payload.start_time or "")
                    row_time = (start + timedelta(milliseconds=(payload.interval or 1) * index)).isoformat(timespec="milliseconds")
                yield name, row_time, value

    @classmethod
    def history_csv(cls, payload: HistoryQueryIn, result: dict[str, Any]) -> bytes:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["point", "time", "value"])
        row_count = 0
        for row in cls._history_rows(payload, result):
            writer.writerow([cls._safe_export_value(value) for value in row])
            row_count += 1
            if row_count > cls.MAX_EXPORT_ROWS:
                raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="导出行数超过安全上限，请缩小范围或增大间隔")
        return ("\ufeff" + output.getvalue()).encode("utf-8")

    @classmethod
    def history_xlsx(cls, payload: HistoryQueryIn, result: dict[str, Any]) -> bytes:
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("History")
        sheet.append(["point", "time", "value"])
        row_count = 0
        xlsx_limit = min(cls.MAX_EXPORT_ROWS, 1_048_575)
        for row in cls._history_rows(payload, result):
            sheet.append([cls._safe_export_value(value) for value in row])
            row_count += 1
            if row_count > xlsx_limit:
                raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Excel 导出超过单工作表行数上限，请改用 CSV 或缩小范围")
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    @classmethod
    def alarm_csv(cls, result: dict[str, Any]) -> bytes:
        fields = ["time", "name", "group", "level", "prio", "ack", "val", "limit", "eu", "desc", "oper", "cat", "parno", "unit"]
        items = result.get("items", [])
        if len(items) > cls.MAX_EXPORT_ROWS:
            raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="报警导出行数超过安全上限，请缩小时间范围或增加筛选条件")
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for item in items:
            writer.writerow({key: cls._safe_export_value(value) for key, value in item.items()})
        return ("\ufeff" + output.getvalue()).encode("utf-8")

    @classmethod
    def alarm_xlsx(cls, result: dict[str, Any]) -> bytes:
        fields = ["time", "name", "group", "level", "prio", "ack", "val", "limit", "eu", "desc", "oper", "cat", "parno", "unit"]
        items = result.get("items", [])
        xlsx_limit = min(cls.MAX_EXPORT_ROWS, 1_048_575)
        if len(items) > xlsx_limit:
            raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="报警 Excel 导出超过单工作表行数上限，请改用 CSV 或增加筛选条件")
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Alarms")
        sheet.append(fields)
        for item in items:
            sheet.append([cls._safe_export_value(item.get(field)) for field in fields])
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    @staticmethod
    async def bytes_stream(content: bytes, chunk_size: int = 64 * 1024) -> AsyncIterator[bytes]:
        for index in range(0, len(content), chunk_size):
            yield content[index : index + chunk_size]
            await asyncio.sleep(0)
